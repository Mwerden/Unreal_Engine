"""
create_best_data_workbook.py

This script DOES NOT change your original Excel file or monotonic CSV files.
It reads:
  1) Temp_Test_Combo.xlsx
  2) MonotonicFrames/<setting folder>/*.csv

It creates a new clean workbook:
  Best_Data_Output.xlsx

Output workbook includes:
  - Summary: one row per setting/test
  - Mapping: which Excel sheet matched which monotonic folder/run
  - Temp_All: all averaged temp data stacked together
  - Merged_All: monotonic timeline + averaged temps
  - One clean sheet per setting/test

Expected temp workbook layout:
  Test_01, Test_02, ... sheets
  Row 3 has headers
  Column A = Time (s)
  Avg columns: H=BIG Avg, O=MID Avg, V=LITTLE Avg, AC=G3D Avg, AJ=SOC Avg
"""

import os
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

# =========================
# USER SETTINGS
# =========================
TEMP_EXCEL_FILE = "Temp_Test_Combo.xlsx"
MONO_ROOT = "MonotonicFrames"
OUTPUT_EXCEL_FILE = "Best_Data_Output.xlsx"

# Options:
#   "run0"  = use the file ending in _0.csv when possible
#   "first" = use the first CSV file found in each setting folder
#   "best"  = choose the run with the best average FPS after processing
MONO_RUN_MODE = "run0"

# If your temp sheets and monotonic folders do not line up alphabetically,
# put the folder names in exact Test_01, Test_02, ... order here.
# Leave empty to use automatic sorted folder order.
FOLDER_ORDER = []

# Keep temps matched to the most recent temperature reading at or before a frame time.
# Change to "nearest" if you want closest temp reading instead.
MERGE_DIRECTION = "backward"

# =========================
# HELPERS
# =========================

def natural_key(text):
    return [int(x) if x.isdigit() else x.lower() for x in re.split(r"(\d+)", str(text))]


def safe_sheet_name(name, used):
    bad = r"[]:*?/\\"
    for ch in bad:
        name = name.replace(ch, "_")
    name = name[:31]
    base = name
    n = 1
    while name in used:
        suffix = f"_{n}"
        name = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(name)
    return name


def find_test_sheets(excel_file):
    xls = pd.ExcelFile(excel_file)
    sheets = [s for s in xls.sheet_names if re.match(r"Test_\d+", s)]
    return sorted(sheets, key=natural_key)


def get_mono_folders(root):
    root_path = Path(root)
    folders = [p.name for p in root_path.iterdir() if p.is_dir()]
    return sorted(folders, key=natural_key)


def read_temp_sheet(excel_file, sheet_name):
    # Header row is Excel row 3, so pandas header index is 2.
    df = pd.read_excel(excel_file, sheet_name=sheet_name, header=2)

    # Use fixed column positions from your workbook.
    needed_positions = {
        0: "time_s",
        7: "BIG_avg",
        14: "MID_avg",
        21: "LITTLE_avg",
        28: "G3D_avg",
        35: "SOC_avg",
    }

    out = pd.DataFrame()
    for pos, new_name in needed_positions.items():
        if pos < len(df.columns):
            out[new_name] = df.iloc[:, pos]
        else:
            out[new_name] = pd.NA

    out = out.dropna(subset=["time_s"])
    out = out[pd.to_numeric(out["time_s"], errors="coerce").notna()].copy()

    for c in out.columns:
        out[c] = pd.to_numeric(out[c], errors="coerce")

    out = out.dropna(subset=["time_s"])
    out = out.sort_values("time_s")
    return out


def normalize_time_to_seconds(series):
    """Convert a time-like column to seconds using simple scale detection."""
    s = pd.to_numeric(series, errors="coerce")
    s = s.dropna()
    if s.empty:
        return series

    values = pd.to_numeric(series, errors="coerce")
    start = values.dropna().iloc[0]
    values = values - start
    max_val = values.max()

    # Heuristics:
    #   nanoseconds often huge
    #   microseconds often millions
    #   milliseconds often thousands to hundreds of thousands
    #   seconds usually under a few hundred for your tests
    if max_val > 1e8:
        return values / 1e9
    elif max_val > 1e5:
        return values / 1e6
    elif max_val > 1000:
        return values / 1000
    else:
        return values


def process_mono_csv(path):
    df = pd.read_csv(path)
    df = df.dropna(how="all")
    if df.empty:
        raise ValueError("empty monotonic CSV")

    # Make all column names strings.
    df.columns = [str(c).strip() for c in df.columns]

    numeric_cols = []
    for c in df.columns:
        vals = pd.to_numeric(df[c], errors="coerce")
        if vals.notna().sum() > 2:
            numeric_cols.append(c)

    if not numeric_cols:
        raise ValueError("no numeric columns found")

    lower_names = {c: c.lower() for c in numeric_cols}

    # Try to find a timestamp/time column.
    time_candidates = [
        c for c in numeric_cols
        if ("time" in lower_names[c] or "timestamp" in lower_names[c] or "mono" in lower_names[c])
        and not any(word in lower_names[c] for word in ["frame_time", "frametime", "delta", "duration", "fps"])
    ]
    time_col = time_candidates[0] if time_candidates else numeric_cols[0]

    out = pd.DataFrame()
    out["time_s"] = normalize_time_to_seconds(df[time_col])
    out = out.dropna(subset=["time_s"]).copy()
    out = out[out["time_s"] >= 0]

    # Try to find an existing frame-time column.
    frame_candidates = [
        c for c in numeric_cols
        if any(word in lower_names[c] for word in ["frame_time", "frametime", "delta", "duration"])
    ]

    if frame_candidates:
        ft = pd.to_numeric(df.loc[out.index, frame_candidates[0]], errors="coerce")
        # Guess units for frame time.
        med = ft.dropna().median()
        if med < 1:          # seconds
            out["frame_time_ms"] = ft * 1000
        elif med > 1000:     # micro/nano-ish, fallback based on size
            out["frame_time_ms"] = ft / 1000
        else:                # already ms
            out["frame_time_ms"] = ft
    else:
        # Compute frame time from timestamp differences.
        out["frame_time_ms"] = out["time_s"].diff() * 1000

    out = out.dropna(subset=["frame_time_ms"])
    out = out[out["frame_time_ms"] > 0]
    out["fps"] = 1000 / out["frame_time_ms"]
    out = out.sort_values("time_s")
    return out[["time_s", "frame_time_ms", "fps"]]


def choose_mono_file(folder_path):
    files = sorted([p for p in Path(folder_path).glob("*.csv")], key=lambda p: natural_key(p.name))
    if not files:
        return None, None

    if MONO_RUN_MODE == "run0":
        run0 = [p for p in files if re.search(r"_0\.csv$", p.name)]
        chosen = run0[0] if run0 else files[0]
        return chosen, process_mono_csv(chosen)

    if MONO_RUN_MODE == "first":
        return files[0], process_mono_csv(files[0])

    if MONO_RUN_MODE == "best":
        best_path = None
        best_df = None
        best_score = -1
        for p in files:
            try:
                df = process_mono_csv(p)
                score = df["fps"].mean()
                if score > best_score:
                    best_score = score
                    best_path = p
                    best_df = df
            except Exception:
                continue
        return best_path, best_df

    raise ValueError("MONO_RUN_MODE must be run0, first, or best")


def merge_mono_temp(mono_df, temp_df):
    merged = pd.merge_asof(
        mono_df.sort_values("time_s"),
        temp_df.sort_values("time_s"),
        on="time_s",
        direction=MERGE_DIRECTION,
    )
    return merged


def style_workbook(path):
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_gray = Side(style="thin", color="D9E2F3")

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
            cell.border = Border(bottom=thin_gray)

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col[:200]:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 28)

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.00"

    wb.save(path)

# =========================
# MAIN PROGRAM
# =========================

def main():
    if not Path(TEMP_EXCEL_FILE).exists():
        raise FileNotFoundError(f"Could not find {TEMP_EXCEL_FILE}")
    if not Path(MONO_ROOT).exists():
        raise FileNotFoundError(f"Could not find folder {MONO_ROOT}")

    test_sheets = find_test_sheets(TEMP_EXCEL_FILE)
    folders = FOLDER_ORDER if FOLDER_ORDER else get_mono_folders(MONO_ROOT)

    total = min(len(test_sheets), len(folders))
    if total == 0:
        raise RuntimeError("No Test_XX sheets or monotonic folders found.")

    summary_rows = []
    mapping_rows = []
    all_temp_rows = []
    all_merged_rows = []
    per_test_outputs = {}

    for idx in range(total):
        test_sheet = test_sheets[idx]
        folder = folders[idx]
        folder_path = Path(MONO_ROOT) / folder

        print(f"Processing {test_sheet}  <->  {folder}")

        temp_df = read_temp_sheet(TEMP_EXCEL_FILE, test_sheet)
        temp_df.insert(0, "setting", folder)
        temp_df.insert(0, "test_sheet", test_sheet)
        all_temp_rows.append(temp_df)

        mono_path, mono_df = choose_mono_file(folder_path)

        if mono_df is None:
            mapping_rows.append({
                "test_sheet": test_sheet,
                "setting_folder": folder,
                "mono_file_used": "NO CSV FOUND",
                "status": "temps only",
            })
            merged = temp_df.copy()
            merged["frame_time_ms"] = pd.NA
            merged["fps"] = pd.NA
        else:
            raw_temp = temp_df.drop(columns=["test_sheet", "setting"])
            merged = merge_mono_temp(mono_df, raw_temp)
            merged.insert(0, "setting", folder)
            merged.insert(0, "test_sheet", test_sheet)

            mapping_rows.append({
                "test_sheet": test_sheet,
                "setting_folder": folder,
                "mono_file_used": mono_path.name,
                "status": "merged",
            })

        all_merged_rows.append(merged)
        per_test_outputs[folder] = merged

        summary_rows.append({
            "test_sheet": test_sheet,
            "setting": folder,
            "mono_file_used": mono_path.name if mono_path else "NO CSV FOUND",
            "avg_BIG_temp": temp_df["BIG_avg"].mean(),
            "avg_MID_temp": temp_df["MID_avg"].mean(),
            "avg_LITTLE_temp": temp_df["LITTLE_avg"].mean(),
            "avg_G3D_temp": temp_df["G3D_avg"].mean(),
            "avg_SOC_temp": temp_df["SOC_avg"].mean(),
            "max_BIG_temp": temp_df["BIG_avg"].max(),
            "max_MID_temp": temp_df["MID_avg"].max(),
            "max_LITTLE_temp": temp_df["LITTLE_avg"].max(),
            "max_G3D_temp": temp_df["G3D_avg"].max(),
            "max_SOC_temp": temp_df["SOC_avg"].max(),
            "avg_fps": merged["fps"].mean() if "fps" in merged.columns else pd.NA,
            "min_fps": merged["fps"].min() if "fps" in merged.columns else pd.NA,
            "max_frame_time_ms": merged["frame_time_ms"].max() if "frame_time_ms" in merged.columns else pd.NA,
        })

    summary_df = pd.DataFrame(summary_rows)
    mapping_df = pd.DataFrame(mapping_rows)
    temp_all_df = pd.concat(all_temp_rows, ignore_index=True)
    merged_all_df = pd.concat(all_merged_rows, ignore_index=True)

    with pd.ExcelWriter(OUTPUT_EXCEL_FILE, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        mapping_df.to_excel(writer, sheet_name="Mapping", index=False)
        temp_all_df.to_excel(writer, sheet_name="Temp_All", index=False)
        merged_all_df.to_excel(writer, sheet_name="Merged_All", index=False)

        used_names = {"Summary", "Mapping", "Temp_All", "Merged_All"}
        for setting, df in per_test_outputs.items():
            sheet_name = safe_sheet_name(setting, used_names)
            df.to_excel(writer, sheet_name=sheet_name, index=False)

    style_workbook(OUTPUT_EXCEL_FILE)
    print(f"\nDONE. Created: {OUTPUT_EXCEL_FILE}")
    print("Your original files were not changed.")


if __name__ == "__main__":
    main()
