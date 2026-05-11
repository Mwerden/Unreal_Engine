"""
clean_best_data_output.py

Purpose:
- Opens your generated Best_Data_Output.xlsx
- Fixes bad FPS rows caused by frame-time unit mismatches
- Rebuilds avg_fps, min_fps, and max_frame_time_ms from the actual merged sheets
- Creates a NEW workbook so your original Excel files are not changed

Expected files in the same folder:
- Best_Data_Output.xlsx

Output:
- Best_Data_Output_CLEANED.xlsx
"""

from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# =========================
# USER SETTINGS
# =========================
INPUT_FILE = "Best_Data_Output.xlsx"
OUTPUT_FILE = "Best_Data_Output_CLEANED.xlsx"

# FPS limits used to detect nonsense values
MAX_REASONABLE_FPS = 300
MIN_REASONABLE_FPS = 1


# =========================
# HELPER FUNCTIONS
# =========================
def find_col(df, possible_names):
    """Find a column using several possible names."""
    lowered = {str(c).strip().lower(): c for c in df.columns}
    for name in possible_names:
        key = name.strip().lower()
        if key in lowered:
            return lowered[key]
    return None


def fix_frame_time_and_fps(df):
    """
    Fixes frame_time/fps units.

    Handles:
    - frame time stored in seconds, like 0.016
    - frame time stored in milliseconds, like 16.0
    - fps accidentally calculated as 1000 / seconds
    """
    df = df.copy()

    frame_col = find_col(df, ["frame_time_ms", "frame_time", "frametime", "frame_ms"])
    fps_col = find_col(df, ["fps"])

    if frame_col is None:
        return df, None

    # Force numeric
    df[frame_col] = pd.to_numeric(df[frame_col], errors="coerce")

    # If mean frame time is less than 1, it is almost certainly seconds
    good_frame_values = df[frame_col].dropna()
    if len(good_frame_values) == 0:
        return df, frame_col

    mean_frame = good_frame_values.mean()

    if mean_frame < 1:
        # Convert seconds to ms
        df["frame_time_ms_clean"] = df[frame_col] * 1000
    else:
        # Already ms
        df["frame_time_ms_clean"] = df[frame_col]

    # Recalculate FPS correctly
    df["fps_clean"] = 1000 / df["frame_time_ms_clean"]

    # Remove impossible values caused by bad rows or bad samples
    df.loc[df["fps_clean"] > MAX_REASONABLE_FPS, "fps_clean"] = pd.NA
    df.loc[df["fps_clean"] < MIN_REASONABLE_FPS, "fps_clean"] = pd.NA

    # Replace or create clean columns
    df["frame_time_ms"] = df["frame_time_ms_clean"]
    df["fps"] = df["fps_clean"]

    # Drop helper columns
    df = df.drop(columns=["frame_time_ms_clean", "fps_clean"])

    return df, "frame_time_ms"


def clean_sheet_name(name):
    """Excel sheet names must be <=31 chars and cannot contain some characters."""
    bad_chars = ["\\", "/", "*", "?", ":", "[", "]"]
    for ch in bad_chars:
        name = name.replace(ch, "_")
    return name[:31]


def auto_width(ws):
    """Make Excel columns easier to read."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            value = cell.value
            if value is not None:
                max_len = max(max_len, len(str(value)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 35)


def style_workbook(path):
    wb = load_workbook(path)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True)
    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
            cell.border = border

        for row in ws.iter_rows():
            for cell in row:
                cell.border = border

        auto_width(ws)

    wb.save(path)


# =========================
# MAIN PROGRAM
# =========================
def main():
    input_path = Path(INPUT_FILE)

    if not input_path.exists():
        raise FileNotFoundError(
            f"Could not find {INPUT_FILE}. Put this script in the same folder as {INPUT_FILE}."
        )

    print(f"Reading {INPUT_FILE}...")

    excel = pd.ExcelFile(input_path)
    output_path = Path(OUTPUT_FILE)

    cleaned_sheets = {}
    summary_rows = []

    for sheet in excel.sheet_names:
        df = pd.read_excel(input_path, sheet_name=sheet)

        # Keep summary/mapping/temp sheets, but clean any sheet that has frame data
        fixed_df, frame_col = fix_frame_time_and_fps(df)

        cleaned_sheets[sheet] = fixed_df

        # Build new corrected summary from merged/test sheets
        if frame_col is not None and "fps" in fixed_df.columns:
            setting_name = sheet

            # Try to use real setting column if it exists
            setting_col = find_col(fixed_df, ["setting"])
            if setting_col is not None and fixed_df[setting_col].dropna().size > 0:
                setting_name = fixed_df[setting_col].dropna().iloc[0]

            mono_col = find_col(fixed_df, ["mono_file_used", "mono_file"])
            mono_file = ""
            if mono_col is not None and fixed_df[mono_col].dropna().size > 0:
                mono_file = fixed_df[mono_col].dropna().iloc[0]

            fps = pd.to_numeric(fixed_df["fps"], errors="coerce")
            frame_time = pd.to_numeric(fixed_df["frame_time_ms"], errors="coerce")

            row = {
                "sheet": sheet,
                "setting": setting_name,
                "mono_file_used": mono_file,
                "avg_fps_clean": round(fps.mean(skipna=True), 2),
                "min_fps_clean": round(fps.min(skipna=True), 2),
                "max_frame_time_ms_clean": round(frame_time.max(skipna=True), 2),
                "bad_fps_rows_removed": int(fixed_df["fps"].isna().sum()),
            }

            # Add temp averages if they exist
            for temp_name in ["BIG", "MID", "LITTLE", "G3D", "SOC"]:
                temp_col = find_col(fixed_df, [temp_name, f"{temp_name}_avg", f"avg_{temp_name}_temp"])
                if temp_col is not None:
                    vals = pd.to_numeric(fixed_df[temp_col], errors="coerce")
                    row[f"avg_{temp_name}_temp"] = round(vals.mean(skipna=True), 2)
                    row[f"max_{temp_name}_temp"] = round(vals.max(skipna=True), 2)

            summary_rows.append(row)

    clean_summary = pd.DataFrame(summary_rows)

    print(f"Writing {OUTPUT_FILE}...")

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        if not clean_summary.empty:
            clean_summary.to_excel(writer, sheet_name="Clean_Summary", index=False)

        for sheet, df in cleaned_sheets.items():
            safe_name = clean_sheet_name(sheet)
            # Avoid overwriting Clean_Summary
            if safe_name == "Clean_Summary":
                safe_name = "Original_Summary"
            df.to_excel(writer, sheet_name=safe_name, index=False)

    style_workbook(output_path)

    print("\nDONE")
    print(f"Created: {OUTPUT_FILE}")
    print("Your original Excel files were not changed.")


if __name__ == "__main__":
    main()
